#!/usr/bin/env python3
"""Published F23 graph data and declared central-value audit; no experimental refit."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import math
import openpyxl
import numpy as np

ROOT=Path(__file__).resolve().parents[2]
RUN=ROOT/'outputs/learning-daily/20261019-DAY17-linear-polarization-paq-compton-run-01'
TMP=ROOT/'tmp/day17-fougeres2023'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    c=json.loads((RUN/'fougeres2023-exercise-contract.json').read_text())
    assert c['identity']=='F23-REAL-SOURCE-DATA-CONTRACT-01'
    for p,h in c['hashes'].items():
        assert sha(ROOT/p)==h
    wb=openpyxl.load_workbook(TMP/'data_figures_article.xlsx',read_only=True,data_only=True)
    sheet=wb['figure 2 right'];profiles=[]
    for label,start,cols in [('5292',6,[5,6,7,8]),('7785',11,[10,11,14,15]),('3796',18,[17,18,19,20])]:
        rows=[]
        for i,row in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
            if row[start-1] is not None:
                rows.append({'excel_row':i,'values':[float(row[k]) for k in cols]})
        a=np.array([row['values'] for row in rows]);x,m,y,e=a.T
        profiles.append({'state_label':label,'sheet':'figure 2 right','excel_columns_1based':[k+1 for k in cols],
                         'n':len(a),'sum_model_displayed':float(sum(m)),'sum_measurement_displayed':float(sum(y)),
                         'max_abs_error_squared_minus_measurement':float(max(abs(e*e-y))),
                         'Pearson_on_all_displayed_bins_control':float(sum((y-m)**2/m)),
                         'Neyman_on_all_displayed_bins_control':float(sum((y-m)**2/e**2)),
                         'actual_rows':rows,
                         'boundary':'displayed graph values only; original fit window/model normalization/covariance not recovered'})
    wb=openpyxl.load_workbook(TMP/'data_figures_supplementary.xlsx',read_only=True,data_only=True)
    grids=[]
    for title in ['figure 1 (a)','figure 1 (b)']:
        rows=[]
        for i,row in enumerate(wb[title].iter_rows(min_row=2,values_only=True),2):
            if all(v is not None for v in row[4:7]):rows.append([i,*[float(v) for v in row[4:7]]])
        a=np.array(rows);positive=a[a[:,3]>0];zero=a[a[:,3]==0]
        grids.append({'sheet':title,'n':len(a),'grid_dimensions':[len(set(a[:,1])),len(set(a[:,2]))],
                      'min_all_row_tau_energy_chi2':a[a[:,3].argmin()].tolist(),
                      'min_positive_row_tau_energy_chi2':positive[positive[:,3].argmin()].tolist(),
                      'zero_points':len(zero),'zero_tau_slices':sorted(set(zero[:,1].tolist())),
                      'boundary':'zero slice retained, not interpreted as optimum; no reconstructed coverage or fit'})
    width=lambda tau:658.2119569/tau
    omega=lambda tau,b,j:(2*j+1)/14*width(tau)*b*(1-b)
    controls=[]
    for tau,b,j,label in [(11,.0068,3.5,'new F23 only'),(10.2,.0066,3.5,'F23 recommended inputs'),(10,.037,3.5,'old branch condition'),(10.2,.0066,1.5,'alternative J3p2 at fixed values'),(10.2,.0066,2.5,'alternative J5p2 at fixed values')]:
        controls.append({'label':label,'tau_fs':tau,'Bp':b,'J':j,'omega_gamma_meV_conditional':omega(tau,b,j)})
    weights=[1/.17**2,1/.08**2]
    br_average=(.68*weights[0]+.65*weights[1])/sum(weights)
    br_scale=1/math.sqrt(sum(weights))
    inversion=[]
    for beta in [.03,.075,.1]:
        for theta in [120,150,170]:
            co=math.cos(math.radians(theta));r=math.sqrt(1-beta*beta)/(1-beta*co)
            recovered=(r*r*co+math.sqrt(1+r*r*co*co-r*r))/(1+r*r*co*co)
            assert abs(recovered-beta)<1e-14
            inversion.append({'beta':beta,'theta_deg':theta,'R':r,'recovered_beta':recovered})
    result={'at':datetime.now(timezone.utc).isoformat(),'script_sha256':sha(Path(__file__)),
            'contract_sha256':sha(RUN/'fougeres2023-exercise-contract.json'),'profiles':profiles,'SI_grids':grids,
            'Table1_quadrature_from_printed_components_pct':math.sqrt(sum(x*x for x in [15,12,12,10,6.5,6,3.5,.3])),
            'Table1_reported_pct':26,'conditional_rate_centers':controls,
            'BR_inverse_variance_control_if_independent_printed_scales':{'mean_percent':br_average,'scale_percent':br_scale},
            'Eq3_4_point_inversion_control':inversion,
            'graph_level_Pearson_not_equal_published_grid_minimum':True,
            'original_fit_not_replayed':True,'no_new_experimental_tau_Bp_spin_or_CI':True,'ok':True}
    (RUN/'fougeres2023-exercise-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in result.items() if k not in ['profiles','Eq3_4_point_inversion_control']},ensure_ascii=False,indent=2))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
